"""Read resumes / job descriptions from PDF, DOCX, TXT or raw bytes."""
from __future__ import annotations

import io
import re


class DocumentError(Exception):
    """Raised when a document cannot be read."""


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\u2022", "- ").replace("\uf0b7", "- ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise DocumentError("Install 'pypdf' to read PDF files.") from e
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        pages = [(p.extract_text() or "") for p in reader.pages]
    except Exception as e:
        raise DocumentError(f"Could not read PDF: {e}") from e
    text = "\n".join(pages)
    if len(text.strip()) < 30:
        raise DocumentError(
            "No text found in this PDF. It may be a scanned image - "
            "export a text-based PDF or paste the text instead."
        )
    return text


def _read_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise DocumentError("Install 'python-docx' to read DOCX files.") from e
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise DocumentError(f"Could not read DOCX: {e}") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def read_document(filename: str, data: bytes) -> str:
    """Return cleaned text for an uploaded file."""
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        text = _read_pdf(data)
    elif name.endswith(".docx"):
        text = _read_docx(data)
    else:
        text = data.decode("utf-8", errors="ignore")
    text = clean_text(text)
    if not text:
        raise DocumentError("The document appears to be empty.")
    return text
